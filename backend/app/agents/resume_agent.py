"""Resume Agent: text extraction, structured parsing and ATS-style scoring."""
import io
import re
from .catalog import ROLES, extract_skills, normalize_skill, resolve_role

MAX_BYTES = 5 * 1024 * 1024

HEADINGS = {
    "summary": r"(professional\s+)?(summary|objective|profile|about\s+me|career\s+objective)",
    "education": r"(education|academic(s| background| qualifications?)?|qualifications?)",
    "skills": r"((technical|key|core)\s+)?skills(\s*(&|and)\s*tools)?|technologies|tech\s+stack|competencies",
    "projects": r"(academic\s+|personal\s+|key\s+)?projects?",
    "experience": r"((work|professional|industrial)\s+)?(experience|internships?|employment)(\s+history)?",
    "certifications": r"(certifications?|certificates?|courses|licen[cs]es?)",
    "achievements": r"(achievements?|awards?(\s*(&|and)\s*honou?rs)?|accomplishments?|honou?rs|extra[- ]?curriculars?|positions? of responsibility)",
    "interests": r"(interests|hobbies)",
}
_HEAD_RE = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in HEADINGS.items()}

ACTION_VERBS = {
    "built", "developed", "designed", "implemented", "created", "led", "optimized", "optimised", "improved", "reduced",
    "increased", "automated", "deployed", "engineered", "integrated", "architected", "launched", "managed", "analyzed",
    "analysed", "trained", "migrated", "collaborated", "delivered", "achieved", "organized", "organised", "mentored",
    "researched", "tested", "refactored", "established", "streamlined", "published", "won", "secured",
}
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)|(?<!\d)\+?\d{1,3}[\s-]?\(?\d{3}\)?[\s-]?\d{3}[\s-]?\d{4}(?!\d)")
LINKS = {
    "linkedin": re.compile(r"linkedin\.com/in/([A-Za-z0-9\-_%]+)", re.I),
    "github": re.compile(r"github\.com/([A-Za-z0-9\-_]+)", re.I),
    "leetcode": re.compile(r"leetcode\.com/(?:u/)?([A-Za-z0-9\-_]+)", re.I),
    "codeforces": re.compile(r"codeforces\.com/profile/([A-Za-z0-9\-_.]+)", re.I),
}
DEGREE = re.compile(r"\b(B\.?\s?E\.?|B\.?\s?Tech|M\.?\s?Tech|B\.?\s?Sc\.?|M\.?\s?Sc\.?|BCA|MCA|MBA|B\.?\s?Com|Bachelor[^\n,|]{0,40}|Master[^\n,|]{0,40})\b", re.I)
BULLET = re.compile(r"^\s*(?:[-•●▪■◦*–]|\d+[.)])\s+")


# ---------------------------------------------------------------- extraction
def extract_text(filename: str, data: bytes) -> str:
    if not data: raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_BYTES: raise ValueError("File too large (max 5 MB).")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        try: reader = PdfReader(io.BytesIO(data))
        except Exception: raise ValueError("Could not read this PDF. Is it corrupted or password-protected?")
        if getattr(reader, "is_encrypted", False):
            try: reader.decrypt("")
            except Exception: raise ValueError("This PDF is password-protected.")
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages[:6])
        if len(text.strip()) < 40:
            raise ValueError("No selectable text found. This looks like a scanned/image PDF; export a text-based PDF or DOCX.")
    elif name.endswith(".docx"):
        import docx
        try: d = docx.Document(io.BytesIO(data))
        except Exception: raise ValueError("Could not read this DOCX file.")
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")
    text = text.replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ------------------------------------------------------------------- parsing
def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    cur = "header"
    for raw in text.split("\n"):
        line = raw.strip()
        if not line: continue
        key = None
        if len(line.split()) <= 5:
            clean = re.sub(r"[^A-Za-z&/ ]", "", line).strip()
            for k, rx in _HEAD_RE.items():
                if rx.match(clean): key = k; break
        if key: cur = key; sections.setdefault(cur, []); continue
        sections.setdefault(cur, []).append(line)
    return sections


def bullets(lines: list[str]) -> list[str]:
    return [BULLET.sub("", l).strip() for l in lines if BULLET.match(l)]


def parse_resume(text: str) -> dict:
    sec = split_sections(text)
    head = "\n".join(sec.get("header", [])[:8])
    em = EMAIL.search(text); ph = PHONE.search(head + "\n" + text[:600])
    contact = {"email": em.group(0) if em else None, "phone": ph.group(0).strip() if ph else None}
    links = {}
    for k, rx in LINKS.items():
        m = rx.search(text)
        if m and m.group(1).lower() not in ("in", "profile", "u"): links[k] = m.group(1)
    name = ""
    for ln in sec.get("header", [])[:4]:
        if "@" in ln or re.search(r"\d{5,}", ln) or len(ln.split()) > 5: continue
        if re.match(r"^[A-Za-z][A-Za-z .'-]{2,}$", ln): name = ln.title() if ln.isupper() else ln; break

    edu_text = "\n".join(sec.get("education", []))
    cgpa = None
    m = re.search(r"(?:cgpa|gpa|c\.g\.p\.a)\s*[:\-]?\s*(\d{1,2}(?:\.\d{1,2})?)(?:\s*/\s*(10|4))?", text, re.I) or \
        re.search(r"(\d\.\d{1,2})\s*/\s*10\b", text)
    if m:
        v = float(m.group(1))
        if len(m.groups()) > 1 and m.groups()[1] == "4": v = round(v / 4 * 10, 2)
        if 0 < v <= 10: cgpa = v
    years = [int(y) for y in re.findall(r"\b(20[12]\d|203[0-2])\b", edu_text)]
    dm = DEGREE.search(edu_text)
    education = {"degree": dm.group(0).strip() if dm else "", "cgpa": cgpa, "graduation_year": max(years) if years else None,
                 "details": sec.get("education", [])[:4]}

    declared = extract_skills("\n".join(sec.get("skills", [])))
    all_found = extract_skills(text)
    projects = _entries(sec.get("projects", []))
    experience = _entries(sec.get("experience", []))
    return {
        "name": name, "contact": contact, "links": links, "education": education,
        "skills": sorted(set(declared) | set(all_found)), "declared_skills": declared,
        "projects": projects, "experience": experience,
        "certifications": [BULLET.sub("", l).strip() for l in sec.get("certifications", [])][:10],
        "achievements": [BULLET.sub("", l).strip() for l in sec.get("achievements", [])][:10],
        "summary": " ".join(sec.get("summary", []))[:500],
        "sections_found": [k for k in HEADINGS if sec.get(k)],
        "word_count": len(text.split()),
    }


def _entries(lines: list[str]) -> list[dict]:
    """Entries = a title line followed by bullet lines."""
    out, cur = [], None
    for ln in lines:
        if BULLET.match(ln):
            if cur is None: cur = {"title": "", "bullets": []}; out.append(cur)
            cur["bullets"].append(BULLET.sub("", ln).strip())
        else:
            if cur is None or cur["bullets"] or len(cur["title"]) > 90:
                cur = {"title": ln, "bullets": []}; out.append(cur)
            else:
                cur["title"] += " " + ln
    return out[:10]


# ------------------------------------------------------------------- scoring
def score_resume(parsed: dict, text: str, target_role: str | None = None) -> dict:
    role = resolve_role(target_role or "")
    br, tips = [], []

    def add(name, got, mx, note):
        br.append({"name": name, "score": round(got, 1), "max": mx, "note": note})

    c = parsed["contact"]; links = parsed["links"]
    got = (4 if c["email"] else 0) + (3 if c["phone"] else 0) + (3 if links.get("linkedin") or links.get("github") else 0)
    add("Contact & links", got, 10, "Email, phone and a LinkedIn/GitHub link")
    if not c["email"]: tips.append("Add a professional email address at the top.")
    if not c["phone"]: tips.append("Add a phone number.")
    if not (links.get("linkedin") or links.get("github")): tips.append("Add your LinkedIn and GitHub URLs.")

    need = ["education", "skills", "projects"]
    have = [s for s in need if s in parsed["sections_found"]]
    got = len(have) * 6 + (4 if parsed["experience"] or "experience" in parsed["sections_found"] else 0) + (3 if parsed["summary"] else 0)
    add("Standard sections", min(25, got + (4 if parsed["certifications"] or parsed["achievements"] else 0)), 25,
        "Education, Skills, Projects, Experience/Internships, Summary")
    for s in need:
        if s not in have: tips.append(f"Add a clearly labelled '{s.title()}' section; ATS parsers look for it.")
    if not parsed["experience"]: tips.append("Add internships, open-source work or freelance work under 'Experience' if you have any.")

    all_bul = [b for e in parsed["projects"] + parsed["experience"] for b in e["bullets"]]
    verbs = sum(1 for b in all_bul if b.split(" ")[0].lower().rstrip(",") in ACTION_VERBS)
    got = min(20, len(all_bul) * 2) * 0.5 + (10 * verbs / len(all_bul) if all_bul else 0)
    add("Bullet quality", got, 20, f"{len(all_bul)} bullets, {verbs} start with action verbs")
    if len(all_bul) < 6: tips.append("Describe each project/experience with 2-3 bullet points.")
    if all_bul and verbs / len(all_bul) < 0.6: tips.append("Start bullets with strong action verbs (Built, Designed, Reduced, Deployed).")

    quant = sum(1 for b in all_bul if re.search(r"\d+\s*(%|x|\+|k\b|ms|users|students|requests|records|hours|days|sec)|\b\d{2,}\b", b, re.I))
    got = min(15, quant * 5)
    add("Measurable impact", got, 15, f"{quant} bullets include numbers")
    if quant < 3: tips.append("Quantify impact (e.g. 'reduced load time by 30%', 'handled 500+ records').")

    wc = parsed["word_count"]
    got = 10 if 250 <= wc <= 900 else 6 if 150 <= wc < 250 or 900 < wc <= 1200 else 2
    add("Length", got, 10, f"{wc} words (ideal 250-900 for a one/two-page CV)")
    if wc < 250: tips.append("Your resume is short; add detail to projects and skills.")
    if wc > 900: tips.append("Trim to one page where possible; keep the strongest content.")

    req = ROLES[role]["core"] + ROLES[role]["nice"]
    mine = set(parsed["skills"])
    hit = sum(1 for r in req if any(x in mine for x in r.split("|")))
    got = 20 * hit / len(req)
    add(f"Keywords for {role}", got, 20, f"{hit} of {len(req)} role keywords present")
    miss = [r.split("|")[0] for r in ROLES[role]["core"] if not any(x in mine for x in r.split("|"))]
    if miss: tips.append(f"For {role}, consider adding (if you know them): {', '.join(miss[:5])}.")

    total = round(sum(b["score"] for b in br))
    return {"ats_score": min(100, total), "target_role": role, "breakdown": br, "suggestions": tips[:8]}


def analyze(filename: str, data: bytes, target_role: str | None = None) -> dict:
    text = extract_text(filename, data)
    parsed = parse_resume(text)
    return {"text": text, "parsed": parsed, **score_resume(parsed, text, target_role)}


def merge_into_profile(profile: dict, parsed: dict) -> dict:
    """Return a profile update that ADDS parsed information without overwriting what the student entered."""
    out = {}
    have = {normalize_skill(s.get("name", "")) for s in profile.get("skills", [])}
    new = [{"name": s, "level": "Intermediate"} for s in parsed["skills"] if s not in have]
    if new: out["skills"] = list(profile.get("skills", [])) + new
    ed = parsed["education"]
    if ed.get("cgpa") and not profile.get("cgpa"): out["cgpa"] = ed["cgpa"]
    if ed.get("graduation_year") and not profile.get("graduation_year"): out["graduation_year"] = ed["graduation_year"]
    if ed.get("degree") and not profile.get("degree"): out["degree"] = ed["degree"]
    if parsed["projects"] and not profile.get("projects"):
        out["projects"] = [{"title": p["title"], "description": " ".join(p["bullets"])[:400]} for p in parsed["projects"][:6]]
    if parsed["experience"] and not profile.get("experiences"):
        out["experiences"] = [{"role": e["title"], "description": " ".join(e["bullets"])[:400]} for e in parsed["experience"][:5]]
    if parsed["certifications"] and not profile.get("certifications"):
        out["certifications"] = [{"name": c} for c in parsed["certifications"][:6]]
    if parsed["achievements"] and not profile.get("achievements"):
        out["achievements"] = parsed["achievements"][:6]
    return out
